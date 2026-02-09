"""
Genera documento Word con los resultados del modelado para la tesis.
"""
import pandas as pd
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# ── Paths ──
RUN_DIR = Path("outputs/runs/run_20260205_183321")
OUT_DOCX = Path("outputs") / "Resultados_Modelado_Tesis.docx"

# ── Data ──
tesis = pd.read_csv(RUN_DIR / "tables" / "tabla_resumen_tesis.csv")
compare = pd.read_csv(RUN_DIR / "tables" / "metrics_compare_all_models.csv")
ordinal_metrics = pd.read_csv(RUN_DIR / "ordinal" / "metrics.csv")
ordinal_by_seed = pd.read_csv(RUN_DIR / "ordinal" / "metrics_by_seed.csv")
sklearn_by_seed = pd.read_csv(RUN_DIR / "tables" / "metrics_by_seed_20260205_183321.csv")

# RF feature importance
fi_path = list((RUN_DIR / "tables").glob("rf_feature_importance_*.csv"))
fi = pd.read_csv(fi_path[0]) if fi_path else None

# ── Helpers ──
def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    shading_elm = shading.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): color_hex,
    })
    shading.append(shading_elm)


def add_styled_table(doc, headers, rows, col_widths=None, highlight_best_col=None):
    """Agrega una tabla con formato profesional."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "2E75B6")

    # Rows
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i % 2 == 0:
                set_cell_shading(cell, "D6E4F0")

    # Column widths
    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Cm(w)

    return table


def fmt(val, decimals=3):
    if pd.isna(val):
        return "—"
    try:
        return f"{float(val):.{decimals}f}"
    except (ValueError, TypeError):
        return str(val)


# ══════════════════════════════════════════════
# BUILD DOCUMENT
# ══════════════════════════════════════════════
doc = Document()

# -- Estilos base --
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# ═══════ PORTADA ═══════
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("RESULTADOS DEL MODELADO PREDICTIVO")
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run(
    "Nivel de Confianza en el Uso de IA en Campañas Políticas Digitales"
)
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(0x59, 0x56, 0x59)

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Run ID: ").bold = True
meta.add_run("20260205_183321\n")
meta.add_run("Fecha: ").bold = True
meta.add_run("5 de febrero de 2026\n")
meta.add_run("Split: ").bold = True
meta.add_run("70% entrenamiento / 30% validación\n")
meta.add_run("Validación cruzada: ").bold = True
meta.add_run("Stratified k-fold (k=5), 5 semillas\n")
meta.add_run("N total: ").bold = True
meta.add_run("316 encuestas válidas")

doc.add_page_break()

# ═══════ 1. CONFIGURACIÓN ═══════
doc.add_heading("1. Configuración Metodológica", level=1)

doc.add_heading("1.1 Variable Dependiente", level=2)
doc.add_paragraph(
    "Nivel de confianza en el uso de IA en campañas políticas digitales (variable ordinal). "
    "Se construye a partir de dos ítems Likert (limpieza del proceso electoral y percepción de fraude) "
    "y se agrupa en 3 clases: Baja (1), Media (2) y Alta (3)."
)

doc.add_heading("1.2 Variables Independientes", level=2)
vars_list = [
    "Conocimiento sobre IA (conoce_ia, identifica_ia, percibe_automatización)",
    "Frecuencia de exposición a contenido político digital",
    "Edad, Género",
    "Facultad / Unidad académica",
    "Rol institucional (estudiante, docente, administrativo)",
    "Percepciones Likert (6 ítems sobre IA en campañas)",
]
for v in vars_list:
    doc.add_paragraph(v, style="List Bullet")

doc.add_heading("1.3 Modelos Evaluados", level=2)
models_desc = [
    ("Random Forest", "Ensemble de árboles de decisión con GridSearchCV (72 combinaciones × 5 folds). "
     "Busca patrones no lineales e interacciones entre variables."),
    ("SVM-RBF", "Máquina de Vectores de Soporte con kernel gaussiano (RBF). "
     "GridSearchCV con 16 combinaciones × 5 folds. Eficaz en espacios de alta dimensión."),
    ("Regresión Logística Ordinal", "Modelo cumulative link (logit/probit) de statsmodels. "
     "CV interna k=5 selecciona entre distribución logit y probit. Respeta la naturaleza ordinal del target."),
    ("Baselines", "Majority (predice clase más frecuente) y Stratified (predice proporcionalmente). "
     "Umbral mínimo que todo modelo debe superar."),
]
for name, desc in models_desc:
    p = doc.add_paragraph()
    p.add_run(f"{name}: ").bold = True
    p.add_run(desc)

doc.add_page_break()

# ═══════ 2. DISTRIBUCIÓN DE CLASES ═══════
doc.add_heading("2. Distribución de Clases", level=1)

doc.add_paragraph(
    "La variable dependiente presenta un desbalance moderado. "
    "La estratificación en los splits garantiza proporciones consistentes."
)

class_data = [
    ["1 — Baja", "117", "37.0%", "82", "35"],
    ["2 — Media", "115", "36.4%", "80", "35"],
    ["3 — Alta", "84", "26.6%", "59", "25"],
    ["Total", "316", "100%", "221", "95"],
]
add_styled_table(
    doc,
    ["Clase", "N Total", "Porcentaje", "Train (70%)", "Test (30%)"],
    class_data,
    col_widths=[3, 2.5, 2.5, 2.5, 2.5],
)

doc.add_paragraph()
doc.add_paragraph(
    "Nota: La clase Baja agrupa los niveles 1 y 2 de confianza, Media el nivel 3, "
    "y Alta los niveles 4 y 5 de la escala original de 5 puntos."
)

doc.add_page_break()

# ═══════ 3. TABLA PRINCIPAL DE RESULTADOS ═══════
doc.add_heading("3. Resultados Comparativos", level=1)

doc.add_heading("3.1 Tabla Resumen (3 clases)", level=2)

model_labels = {
    "random_forest": "Random Forest",
    "svm_rbf": "SVM-RBF",
    "ordinal_logit": "Reg. Logística Ordinal",
    "baseline_stratified": "Baseline Estratificado",
    "baseline_majority": "Baseline Mayoría",
}

result_rows = []
for _, r in tesis.iterrows():
    if r["n_classes"] == 3:
        label = model_labels.get(r["model"], r["model"])
        result_rows.append([
            label,
            fmt(r["accuracy"]),
            fmt(r["f1_weighted"]),
            fmt(r["kappa_qw"]),
            fmt(r["auc_ovr"]),
            fmt(r["time_seconds"], 1),
        ])

# Add ordinal 5-class
ord5 = tesis[tesis["n_classes"] == 5]
if not ord5.empty:
    r = ord5.iloc[0]
    result_rows.insert(3, [
        "Reg. Log. Ordinal (5 niveles)",
        fmt(r["accuracy"]),
        fmt(r["f1_weighted"]),
        fmt(r["kappa_qw"]),
        fmt(r["auc_ovr"]),
        fmt(r["time_seconds"], 1),
    ])

add_styled_table(
    doc,
    ["Modelo", "Accuracy", "F1-weighted", "Kappa QW", "AUC OvR", "Tiempo (s)"],
    result_rows,
    col_widths=[5, 2, 2.5, 2, 2, 2.5],
)

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Interpretación: ").bold = True
p.add_run(
    "Random Forest y SVM-RBF obtienen resultados prácticamente idénticos en accuracy (0.531) "
    "y F1-weighted (0.529). El RF lidera en AUC (0.711 vs 0.681), mientras que SVM tiene "
    "un Kappa ligeramente superior (0.320 vs 0.313) y es 26x más rápido. "
    "Ambos superan claramente los baselines y la regresión ordinal."
)

doc.add_page_break()

# ═══════ 4. DETALLE POR SEMILLA ═══════
doc.add_heading("3.2 Estabilidad por Semilla (RF y SVM)", level=2)

doc.add_paragraph(
    "Se evaluaron 5 semillas aleatorias [0, 7, 13, 21, 42] para medir la "
    "estabilidad de los modelos ante variaciones en el split de datos."
)

for model_name in ["random_forest", "svm_rbf"]:
    label = model_labels[model_name]
    doc.add_heading(f"{label}", level=3)
    subset = sklearn_by_seed[sklearn_by_seed["model"] == model_name]
    seed_rows = []
    for _, r in subset.iterrows():
        seed_rows.append([
            int(r["seed"]),
            fmt(r["accuracy"]),
            fmt(r["f1_weighted"]),
            fmt(r["kappa_qw"]),
            fmt(r.get("auc_ovr", float("nan"))),
            fmt(r.get("time_seconds", 0), 1),
        ])
    # Add mean row
    run_csv = pd.read_csv(RUN_DIR / "tables" / "metrics_run_20260205_183321.csv")
    run_row = run_csv[run_csv["model"] == model_name].iloc[0]
    seed_rows.append([
        "MEDIA",
        fmt(run_row["accuracy_mean"]),
        fmt(run_row["f1_weighted_mean"]),
        fmt(run_row["kappa_qw_mean"]),
        fmt(run_row.get("auc_ovr_mean", float("nan"))),
        fmt(run_row.get("time_seconds_mean", 0), 1),
    ])

    add_styled_table(
        doc,
        ["Seed", "Accuracy", "F1-w", "Kappa QW", "AUC", "Tiempo (s)"],
        seed_rows,
        col_widths=[2, 2.5, 2.5, 2.5, 2.5, 2.5],
    )
    doc.add_paragraph()

# ═══════ 5. ORDINAL DETALLE ═══════
doc.add_heading("3.3 Regresión Logística Ordinal (detalle por semilla)", level=2)

doc.add_paragraph(
    "Para cada semilla, la validación cruzada interna (k=5) selecciona automáticamente "
    "la distribución (logit o probit) con mejor F1-weighted en el conjunto de entrenamiento."
)

ord_seed_rows = []
for _, r in ordinal_by_seed.iterrows():
    ord_seed_rows.append([
        int(r["seed"]),
        f"{int(r['n_classes'])}c",
        r.get("distribution", "—"),
        fmt(r["accuracy"]),
        fmt(r["f1_weighted"]),
        fmt(r["kappa_qw"]),
        fmt(r.get("auc_ovr", float("nan"))),
        fmt(r.get("time_seconds", 0), 1),
    ])

add_styled_table(
    doc,
    ["Seed", "Clases", "Distr.", "Acc", "F1-w", "Kappa", "AUC", "Tiempo"],
    ord_seed_rows,
    col_widths=[1.5, 1.5, 1.5, 2, 2, 2, 2, 2],
)
doc.add_paragraph()

doc.add_page_break()

# ═══════ 6. COMPARACIÓN CON RUN ANTERIOR ═══════
doc.add_heading("4. Impacto del Cambio de Split (80/20 → 70/30)", level=1)

doc.add_paragraph(
    "Se compararon los resultados del run anterior (80/20, sin AUC ni tiempo) "
    "con el run actual (70/30, metodología de tesis completa)."
)

delta_rows = [
    ["Random Forest", "Accuracy", "0.597", "0.531", "-0.066"],
    ["", "F1-weighted", "0.595", "0.529", "-0.066"],
    ["", "Kappa QW", "0.431", "0.313", "-0.117"],
    ["SVM-RBF", "Accuracy", "0.578", "0.531", "-0.048"],
    ["", "F1-weighted", "0.571", "0.529", "-0.043"],
    ["", "Kappa QW", "0.388", "0.320", "-0.068"],
]

add_styled_table(
    doc,
    ["Modelo", "Métrica", "Old (80/20)", "New (70/30)", "Δ"],
    delta_rows,
    col_widths=[3.5, 2.5, 2.5, 2.5, 2],
)

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Análisis: ").bold = True
p.add_run(
    "La caída de rendimiento es esperada y metodológicamente justificable. Al reducir el "
    "conjunto de entrenamiento de 252 a 221 observaciones (-12.3%) y aumentar el test de 64 a 95 "
    "(+48.4%), la evaluación se vuelve más exigente. El split 70/30 es más conservador y ofrece "
    "una estimación más realista del rendimiento en datos no vistos, alineándose con la "
    "metodología especificada en la tesis."
)

doc.add_page_break()

# ═══════ 7. IMPORTANCIA DE VARIABLES (RF) ═══════
doc.add_heading("5. Importancia de Variables (Random Forest)", level=1)

doc.add_paragraph(
    "Las 15 variables más importantes según la importancia de Gini del Random Forest:"
)

if fi is not None:
    fi_sorted = fi.sort_values("importance", ascending=False).head(15)
    fi_rows = []
    for rank, (_, r) in enumerate(fi_sorted.iterrows(), 1):
        fi_rows.append([
            rank,
            r["feature"][:60] + ("..." if len(str(r["feature"])) > 60 else ""),
            fmt(r["importance"], 4),
        ])

    add_styled_table(
        doc,
        ["#", "Variable", "Importancia (Gini)"],
        fi_rows,
        col_widths=[1, 10, 3],
    )

doc.add_paragraph()

doc.add_page_break()

# ═══════ 8. FIGURAS ═══════
doc.add_heading("6. Visualizaciones", level=1)

figures = [
    ("compare_models.png", "Comparación de F1-weighted y Kappa QW por modelo"),
    ("compare_auc.png", "Comparación de AUC One-vs-Rest por modelo"),
]

for fname, caption in figures:
    fpath = RUN_DIR / "figures" / fname
    if fpath.exists():
        doc.add_heading(caption, level=3)
        doc.add_picture(str(fpath), width=Inches(5.5))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

# Confusion matrices
doc.add_heading("Matrices de Confusión", level=3)
cm_files = sorted((RUN_DIR / "figures").glob("cm_*.png"))
for cm_path in cm_files:
    model_tag = cm_path.stem.replace("cm_", "").split("_2026")[0]
    nice_name = model_tag.replace("_", " ").title()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(nice_name)
    run.italic = True
    run.font.size = Pt(10)
    doc.add_picture(str(cm_path), width=Inches(3.5))
    last_p = doc.paragraphs[-1]
    last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

# RF importance figure
rf_fig = list((RUN_DIR / "figures").glob("rf_feature_importance_*.png"))
if rf_fig:
    doc.add_heading("Importancia de Variables (RF)", level=3)
    doc.add_picture(str(rf_fig[0]), width=Inches(5.5))
    last_p = doc.paragraphs[-1]
    last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Ordinal figures
ord_figs = sorted((RUN_DIR / "figures" / "ordinal").glob("*.png"))
if ord_figs:
    doc.add_page_break()
    doc.add_heading("Figuras del Modelo Ordinal", level=3)
    for fig_path in ord_figs:
        nice = fig_path.stem.replace("_", " ").title()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(nice)
        run.italic = True
        run.font.size = Pt(10)
        doc.add_picture(str(fig_path), width=Inches(4.0))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

# ═══════ 9. CONCLUSIONES ═══════
doc.add_page_break()
doc.add_heading("7. Conclusiones Preliminares", level=1)

conclusions = [
    "Random Forest y SVM-RBF son los mejores modelos con rendimiento estadísticamente "
    "equivalente (F1-w ≈ 0.529, Acc ≈ 0.531).",
    "RF ofrece mejor AUC (0.711 vs 0.681) y mayor interpretabilidad vía importancia de variables.",
    "SVM-RBF es significativamente más rápido (9.9s vs 260.9s) con rendimiento comparable.",
    "La Regresión Logística Ordinal (F1-w = 0.440) muestra rendimiento inferior, sugiriendo "
    "que las relaciones no lineales en los datos favorecen modelos más flexibles.",
    "Todos los modelos superan claramente los baselines (F1-w ≈ 0.20-0.34), "
    "confirmando que existe señal predictiva en las variables.",
    "El cambio de split 80/20 a 70/30 redujo las métricas pero proporciona una evaluación "
    "más robusta y conservadora, alineada con la metodología de la tesis.",
]
for c in conclusions:
    doc.add_paragraph(c, style="List Number")

# ═══════ SAVE ═══════
doc.save(str(OUT_DOCX))
print(f"\nDocumento generado exitosamente: {OUT_DOCX}")
print(f"Tamaño: {OUT_DOCX.stat().st_size / 1024:.0f} KB")
